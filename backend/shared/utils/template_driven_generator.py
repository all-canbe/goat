"""
模板驱动的内容注入器
在模板 docx 上注入 AI 生成的结构化内容，保留模板的样式定义、页面设置和节结构。

退级机制：
- Level 2: 单个元素注入失败时跳过，用占位文本替代
- Level 3: 整体注入崩溃时退级到 Pandoc 路线
"""

import os
import re
import copy
import json
import logging
from typing import List, Optional, Dict, Tuple
from dataclasses import dataclass, field

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from pydantic import BaseModel

logger = logging.getLogger(__name__)


# ──────────────────────────── 数据模型 ────────────────────────────

class RunAnnotation(BaseModel):
    """run 级别的格式标注"""
    start: int
    end: int
    bold: bool = False
    italic: bool = False
    superscript: bool = False
    subscript: bool = False
    color: Optional[List[int]] = None


class ContentBlock(BaseModel):
    """AI 生成的单个内容块"""
    semantic_type: str = "body"
    level: int = 0
    text: str = ""
    runs: List[RunAnnotation] = []
    children: List["ContentBlock"] = []


class TableBlock(BaseModel):
    """表格数据"""
    caption: str = ""
    headers: List[str] = []
    rows: List[List[str]] = []


class ImageBlock(BaseModel):
    """图片数据"""
    path: str = ""
    caption: str = ""
    width_cm: float = 14.0


class StructuredContent(BaseModel):
    """AI 输出的结构化内容"""
    title: str = ""
    blocks: List[ContentBlock] = []
    tables: List[TableBlock] = []
    images: List[ImageBlock] = []


@dataclass
class ElementInjectionWarning:
    """元素注入警告"""
    element_type: str
    element_index: int
    reason: str


@dataclass
class GenerationResult:
    """生成结果"""
    output_path: str = ""
    degradation_level: int = 0
    used_pandoc_fallback: bool = False
    used_default_template: bool = False
    total_failure: bool = False
    warnings: List[str] = field(default_factory=list)
    element_warnings: List[ElementInjectionWarning] = field(default_factory=list)


# ──────────────────────────── 样式名映射 ────────────────────────────

HEADING_LEVEL_PATTERNS: Dict[int, List[str]] = {
    1: ["标题 1", "标题1", "Heading 1", "heading 1", "一级标题"],
    2: ["标题 2", "标题2", "Heading 2", "heading 2", "二级标题"],
    3: ["标题 3", "标题3", "Heading 3", "heading 3", "三级标题"],
    4: ["标题 4", "标题4", "Heading 4", "heading 4", "四级标题"],
    5: ["标题 5", "标题5", "Heading 5", "heading 5", "五级标题"],
    6: ["标题 6", "标题6", "Heading 6", "heading 6", "六级标题"],
}

BODY_STYLE_PATTERNS = ["正文", "Normal", "normal", "Body", "body", "默认段落字体"]

LIST_STYLE_PATTERNS = [
    "List Paragraph", "List Paragraph", "列表段落",
    "List Number", "List Bullet", "List Continue",
]

QUOTE_STYLE_PATTERNS = ["Quote", "quote", "引用", "Block Text"]


def find_style_name(template_doc, semantic_type: str, level: int = 0) -> Tuple[str, str]:
    """
    在模板中查找最匹配的样式名。
    Returns: (style_name, match_quality): "exact" | "fuzzy" | "default"
    """
    available = {s.name for s in template_doc.styles if s.type == 1}

    if semantic_type.startswith("heading"):
        if level == 0:
            m = re.search(r'\d+', semantic_type)
            level = int(m.group()) if m else 1

        # 精确匹配
        for pattern in HEADING_LEVEL_PATTERNS.get(level, []):
            if pattern in available:
                return pattern, "exact"

        # 模糊匹配
        for name in available:
            if re.search(rf'(标题|Heading|TITLE)\s*{level}', name, re.IGNORECASE):
                return name, "fuzzy"

        # outline_level 匹配
        for style in template_doc.styles:
            if style.type == 1 and style.paragraph_format:
                try:
                    if style.paragraph_format.outline_level == level - 1:
                        return style.name, "fuzzy"
                except Exception:
                    pass

        return "Normal", "default"

    elif semantic_type == "body":
        for pattern in BODY_STYLE_PATTERNS:
            if pattern in available:
                return pattern, "exact"
        return "Normal", "default"

    elif semantic_type == "list":
        for pattern in LIST_STYLE_PATTERNS:
            if pattern in available:
                return pattern, "exact"
        return find_style_name(template_doc, "body")[0], "default"

    elif semantic_type == "quote":
        for pattern in QUOTE_STYLE_PATTERNS:
            if pattern in available:
                return pattern, "exact"
        return find_style_name(template_doc, "body")[0], "default"

    elif semantic_type == "caption":
        for name in available:
            if "caption" in name.lower() or "题注" in name.lower():
                return name, "exact"
        return find_style_name(template_doc, "body")[0], "default"

    return "Normal", "default"


# ──────────────────────────── 内容注入器 ────────────────────────────

class TemplateDrivenGenerator:
    """
    模板驱动的文档生成器。
    在模板 docx 上注入 AI 生成的结构化内容。
    """

    def __init__(self, template_doc: Document, registry=None):
        self.template_doc = template_doc
        self.registry = registry

    def generate(self, content: StructuredContent, output_path: str) -> GenerationResult:
        """
        主入口：在模板上注入内容，输出 docx。
        每个元素独立 try-except，失败时跳过并记录警告。
        """
        result = GenerationResult()
        doc = copy.deepcopy(self.template_doc)

        # 清除模板中的示例内容
        self._clear_body_content(doc)

        # 注入标题（如果有）
        if content.title:
            try:
                style_name, quality = find_style_name(doc, "heading", 1)
                if quality != "exact":
                    result.degradation_level = max(result.degradation_level, 1)
                p = doc.add_paragraph(style=style_name)
                p.text = content.title
                _apply_run_annotations(p, [RunAnnotation(
                    start=0, end=len(content.title), bold=True
                )])
            except Exception as e:
                logger.warning(f"标题注入失败: {e}")
                result.element_warnings.append(ElementInjectionWarning(
                    element_type="title", element_index=0, reason=str(e)
                ))
                try:
                    doc.add_paragraph(content.title)
                except Exception:
                    pass

        # 注入正文块
        for i, block in enumerate(content.blocks):
            try:
                self._inject_paragraph(doc, block, result)
            except Exception as e:
                logger.warning(f"段落注入失败 [block#{i}]: {e}")
                result.element_warnings.append(ElementInjectionWarning(
                    element_type=block.semantic_type, element_index=i, reason=str(e)
                ))
                try:
                    doc.add_paragraph(f"[内容插入失败: {block.text[:80]}...]")
                except Exception:
                    pass

        # 注入表格
        for i, table in enumerate(content.tables):
            try:
                self._inject_table(doc, table, result)
            except Exception as e:
                logger.warning(f"表格注入失败 [table#{i}]: {e}")
                result.element_warnings.append(ElementInjectionWarning(
                    element_type="table", element_index=i, reason=str(e)
                ))
                try:
                    doc.add_paragraph(f"[表格插入失败: {table.caption}]")
                except Exception:
                    pass

        # 注入图片
        for i, image in enumerate(content.images):
            try:
                self._inject_image(doc, image, result)
            except Exception as e:
                logger.warning(f"图片注入失败 [image#{i}]: {e}")
                result.element_warnings.append(ElementInjectionWarning(
                    element_type="image", element_index=i, reason=str(e)
                ))
                try:
                    doc.add_paragraph(f"[图片插入失败: {image.caption or image.path}]")
                except Exception:
                    pass

        # 保存
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        doc.save(output_path)
        result.output_path = output_path

        if result.element_warnings:
            result.degradation_level = max(result.degradation_level, 2)

        return result

    # ────────────── 内部方法 ──────────────

    def _clear_body_content(self, doc: Document):
        """删除模板中的段落内容，但保留样式定义、页面设置和页眉页脚。"""
        body = doc.element.body
        children = list(body)
        for child in children:
            tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
            if tag in ('p', 'tbl'):
                body.remove(child)

    def _inject_paragraph(self, doc: Document, block: ContentBlock, result: GenerationResult):
        """注入一个段落。"""
        style_name, quality = find_style_name(doc, block.semantic_type, block.level)
        if quality != "exact":
            result.degradation_level = max(result.degradation_level, 1)

        p = doc.add_paragraph(style=style_name)
        p.text = block.text

        # 应用 run 级别格式标注
        if block.runs:
            _apply_run_annotations(p, block.runs)

    def _inject_table(self, doc: Document, table: TableBlock, result: GenerationResult):
        """注入一个表格。"""
        # 表格标题
        if table.caption:
            try:
                caption_style, _ = find_style_name(doc, "caption")
                cp = doc.add_paragraph(style=caption_style)
                cp.text = table.caption
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                doc.add_paragraph(table.caption)

        if not table.headers and not table.rows:
            return

        rows = len(table.rows) + (1 if table.headers else 0)
        cols = max(len(table.headers), max((len(r) for r in table.rows), default=0))
        if cols == 0:
            return

        word_table = doc.add_table(rows=rows, cols=cols)

        # 尝试设置表格样式
        try:
            if 'Table Grid' in [s.name for s in doc.styles]:
                word_table.style = 'Table Grid'
        except Exception:
            pass

        # 填充表头
        row_idx = 0
        if table.headers:
            for j, header in enumerate(table.headers):
                if j < cols:
                    cell = word_table.rows[row_idx].cells[j]
                    cell.text = header
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
            row_idx += 1

        # 填充数据行
        for row_data in table.rows:
            for j, cell_text in enumerate(row_data):
                if j < cols:
                    word_table.rows[row_idx].cells[j].text = cell_text
            row_idx += 1

        # 基础边框（兜底）
        try:
            self._set_table_borders(word_table)
        except Exception:
            pass

    def _inject_image(self, doc: Document, image: ImageBlock, result: GenerationResult):
        """插入一张图片。"""
        if not image.path or not os.path.exists(image.path):
            doc.add_paragraph(f"[图片未找到: {image.path}]")
            return

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(image.path, width=Cm(image.width_cm))

        # 图片题注
        if image.caption:
            caption_p = doc.add_paragraph()
            caption_p.text = image.caption
            caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                for run in caption_p.runs:
                    run.font.size = Pt(10.5)
            except Exception:
                pass

    def _set_table_borders(self, table):
        """为表格设置基础单线边框。"""
        tbl = table._tbl
        tblPr = tbl.tblPr if tbl.tblPr is not None else tbl._add_tblPr()

        borders = tblPr.find(qn('w:tblBorders'))
        if borders is not None:
            tblPr.remove(borders)

        from docx.oxml.shared import OxmlElement
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            element = OxmlElement(f'w:{edge}')
            element.set(qn('w:val'), 'single')
            element.set(qn('w:sz'), '4')
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), 'auto')
            borders.append(element)
        tblPr.append(borders)


# ──────────────────────────── 辅助函数 ────────────────────────────

def _apply_run_annotations(paragraph, runs: List[RunAnnotation]):
    """对段落中的文本应用 run 级别格式标注。"""
    if not paragraph.runs:
        return

    full_text = paragraph.text
    for run in paragraph.runs:
        run_text = run.text
        run_start = full_text.find(run_text)
        if run_start < 0:
            continue

        for ann in runs:
            if ann.start <= run_start and run_start + len(run_text) <= ann.end:
                if ann.bold:
                    run.font.bold = True
                if ann.italic:
                    run.font.italic = True
                if ann.superscript:
                    run.font.superscript = True
                if ann.subscript:
                    run.font.subscript = True
                if ann.color and len(ann.color) == 3:
                    run.font.color.rgb = RGBColor(*ann.color)
                break


def render_structured_content_to_html(content: StructuredContent) -> str:
    """将结构化内容渲染为 HTML（供 Pandoc 退级路线使用）。"""
    parts = [
        "<!DOCTYPE html>",
        "<html><head><meta charset='utf-8'>",
        "<style>",
        "body { font-family: SimSun, Times New Roman; line-height: 1.5; }",
        "h1 { text-align: center; font-family: SimHei; font-size: 20pt; }",
        "h2 { font-family: SimHei; font-size: 16pt; }",
        "h3 { font-family: SimHei; font-size: 14pt; }",
        "p { text-indent: 2em; text-align: justify; }",
        "table { border-collapse: collapse; width: 100%; }",
        "th, td { border: 1px solid #000; padding: 6pt; }",
        "</style></head><body>",
    ]

    if content.title:
        parts.append(f"<h1>{content.title}</h1>")

    for block in content.blocks:
        if block.semantic_type.startswith("heading"):
            level = block.level or int(re.search(r'\d', block.semantic_type).group() or 1)
            parts.append(f"<h{level}>{block.text}</h{level}>")
        elif block.semantic_type == "list":
            parts.append(f"<p style='text-indent:0'>• {block.text}</p>")
        elif block.semantic_type == "quote":
            parts.append(f"<blockquote><p>{block.text}</p></blockquote>")
        else:
            parts.append(f"<p>{block.text}</p>")

    for table in content.tables:
        if table.caption:
            parts.append(f"<p style='text-align:center;text-indent:0'><b>{table.caption}</b></p>")
        if table.headers:
            parts.append("<table><thead><tr>")
            for h in table.headers:
                parts.append(f"<th>{h}</th>")
            parts.append("</tr></thead><tbody>")
        for row in table.rows:
            parts.append("<tr>")
            for cell in row:
                parts.append(f"<td>{cell}</td>")
            parts.append("</tr>")
        if table.headers:
            parts.append("</tbody></table>")

    parts.append("</body></html>")
    return "\n".join(parts)
